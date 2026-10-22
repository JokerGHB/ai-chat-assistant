# -*- coding: utf-8 -*-
"""
按 C:\\Users\\lenovo\\Desktop\\moba\\AI应用开发工程师.pdf 的版式重建一份 Java 简历。
几何参数来自 pypdf 反解（A4 595.3x841.9pt，左/右 57pt，正文 10.5pt/18.2pt 行距）。
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Microsoft YaHei"
OUT_DIR = r"D:\work\java\yijia-ai-chat-assistant\resume"
OUT = os.path.join(OUT_DIR, "李一甲-Java后端开发工程师.docx")
os.makedirs(OUT_DIR, exist_ok=True)

BODY = 10.5      # 正文字号
SMALL = 9.0      # 眉标/联系方式/技术栈/URL
NAME_SZ = 11.0   # 姓名行 / 项目标题
TITLE = 20.0     # 主标题
HEAD = 12.0      # 小节标题

LINE = 18.2      # 正文行距
LINE_S = 16.2    # 9pt 行距
LINE_T = 25.0    # 主标题行距


def set_run_font(run, size, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), FONT)


def para(doc_or_hf, *, before=0.0, after=0.0, line=LINE, align=WD_ALIGN_PARAGRAPH.LEFT,
         left=0.0, first=0.0, size=BODY):
    p = doc_or_hf.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line)
    pf.alignment = align
    pf.left_indent = Pt(left)
    pf.first_line_indent = Pt(first)
    pf.widow_control = False
    # 关闭对齐网格，保证行距与 PDF 一致
    ppr = p._element.get_or_add_pPr()
    snap = OxmlElement("w:snapToGrid")
    snap.set(qn("w:val"), "0")
    ppr.append(snap)
    return p


def add_runs(p, parts, size=BODY):
    """parts: [(text, bold), ...]"""
    for text, bold in parts:
        r = p.add_run(text)
        set_run_font(r, size, bold)
    return p


def build():
    doc = Document()

    # ---- 默认样式 ----
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(BODY)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.space_after = Pt(0)

    # ---- 页面 ----
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Pt(57)
    sec.right_margin = Pt(57)
    sec.top_margin = Pt(62)
    sec.bottom_margin = Pt(56)
    sec.header_distance = Pt(32)
    sec.footer_distance = Pt(36)

    # ---- 页眉：姓名 · 方向（每页重复） ----
    hdr = sec.header
    hdr.is_linked_to_previous = False
    hp = hdr.paragraphs[0]
    hp.paragraph_format.space_before = Pt(0)
    hp.paragraph_format.space_after = Pt(0)
    hp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    hp.paragraph_format.line_spacing = Pt(LINE_S)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = hp.add_run("李一甲 · Java 后端开发实习生")
    set_run_font(r, SMALL)

    # ---- 页脚：— 第 N 页 — ----
    ftr = sec.footer
    ftr.is_linked_to_previous = False
    fp = ftr.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    fp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fp.paragraph_format.line_spacing = Pt(LINE_S)
    r1 = fp.add_run("— ")
    set_run_font(r1, SMALL)
    # PAGE 域
    run = fp.add_run()
    set_run_font(run, SMALL)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    r2 = fp.add_run(" —")
    set_run_font(r2, SMALL)

    # ================= 头部 =================
    p = para(doc, before=0, line=LINE_T, align=WD_ALIGN_PARAGRAPH.CENTER, size=TITLE)
    add_runs(p, [("Java 后端开发工程师", True)], TITLE)

    p = para(doc, before=4, line=LINE, align=WD_ALIGN_PARAGRAPH.CENTER, size=NAME_SZ)
    add_runs(p, [("姓名：李一甲　　年龄：21", False)], NAME_SZ)

    p = para(doc, before=2, line=LINE_S, align=WD_ALIGN_PARAGRAPH.CENTER, size=SMALL)
    add_runs(p, [("136-1399-9293 | m15038715470@163.com | github.com/JokerGHB", False)], SMALL)

    # ================= 个人简介 =================
    p = para(doc, before=7.5, after=7.0, line=LINE, size=HEAD)
    add_runs(p, [("个人简介", True)], HEAD)

    p = para(doc, line=LINE)
    add_runs(p, [(
        "网络工程大四，熟悉 Java / Spring Boot 后端开发与大模型应用落地。独立完成多用户 AI 对话系统，"
        "覆盖令牌鉴权、数据隔离、管理后台与统计看板全链路；掌握 MyBatis 开发、MySQL 索引与预聚合设计、"
        "BCrypt 密码安全、定时任务与统一异常处理，并具备 RAG、LangGraph Agent 大模型工程经验。",
        False)], BODY)

    # ================= 教育背景 =================
    p = para(doc, before=7.5, after=7.0, line=LINE, size=HEAD)
    add_runs(p, [("教育背景", True)], HEAD)

    p = para(doc, line=LINE)
    add_runs(p, [("南京工程学院", True), (" | 网络工程（本科）", False)])

    p = para(doc, before=3, line=LINE)
    add_runs(p, [("2023.09 — 2027.07", False)], SMALL)
    p.add_run("    ")
    add_runs(p, [("专业排名: 6/36", False)], BODY)

    p = para(doc, before=3, line=LINE)
    add_runs(p, [("主修课程：", True), (
        "Python 程序设计、Linux 编程、C 语言程序设计、Java 语言程序设计、"
        "数据结构与算法、计算机网络、数据库原理、WEB 前端开发技术", False)])

    # ================= 项目经历 =================
    p = para(doc, before=7.5, after=7.0, line=LINE, size=HEAD)
    add_runs(p, [("项目经历", True)], HEAD)

    # ---------- 项目一 ----------
    p = para(doc, before=7.0, line=LINE)
    add_runs(p, [("拾光 AI 聊天系统（多用户版）  ", True), ("2026.10 - 至今", True)], NAME_SZ)

    p = para(doc, line=LINE_S)
    add_runs(p, [("github.com/JokerGHB/ai-chat-assistant", False)], SMALL)

    p = para(doc, before=1, line=LINE_S)
    add_runs(p, [("技术栈：", True), ("Java 17 / Spring Boot 4 / MyBatis / MySQL 8 / Vue 3 / Vite / RESTful", False)], SMALL)

    p = para(doc, before=1.5, line=LINE)
    add_runs(p, [(
        "基于 Spring Boot + Vue 3 独立开发的多用户 AI 对话系统，落地注册登录、会话与消息的用户级隔离、"
        "每日次数限额与敏感词过滤，并提供用户封禁、会话管理、ECharts 统计看板的管理员后台；"
        "后端 56 个类约 2000 行、前端约 1300 行。", False)])

    bullets1 = [
        ("分层架构与多用户数据隔离",
         "采用 Controller / Service / Mapper 三层结构，MyBatis 全注解 Mapper 免去 XML 维护；"
         "会话与消息的读写一律携带 user_id 条件，跨用户访问统一返回 404 而非 403，不泄露他人数据是否存在。"),
        ("服务端令牌鉴权与密码安全",
         "放弃 JWT，登录签发 32 字节 SecureRandom 令牌落库，拦截器逐请求校验并滑动刷新 last_active_at"
         "（闲置 30 分钟失效）；封禁、改密即时删除令牌使权限变更秒级生效。密码统一 BCrypt（随机盐 + 成本因子 10），"
         "登录失败与账号不存在返回同一文案，防止账号枚举。"),
        ("索引设计与预聚合优化",
         "按查询路径规划联合索引，idx_message_user_daily(user_id, role, created_at) 支撑每日限额实时计数、"
         "idx_session_purge(last_active_at) 支撑过期清理；高频提问由全表 GROUP BY 改为 chat_question_stat "
         "预聚合表（trim + 小写取 MD5 指纹去重、ask_count 原子自增），看板“全部”口径从全表扫描降为主键级查询。"),
        ("对话链路编排与友好降级",
         "单轮对话按「敏感词拦截 → 每日限额 → 先落库提问 → 首条自动生成标题 → 调用大模型」顺序编排；"
         "模型侧 401/429/超时统一映射为业务异常，失败时写入 status='failed' 的 assistant 消息并仍返回 200，"
         "用户提问不丢失，管理端可据此统计“未解决对话”。"),
        ("成本核算与可观测性",
         "用 BigDecimal 按 prompt / completion token 单价在消息落库时计算并持久化 cost；看板按 7d / 30d / all "
         "区间聚合调用次数、总成本、未解决数与高频提问 Top N；登录、封禁、重置密码、删除会话、定时清理等"
         "关键操作统一写入 operation_log 审计表。"),
        ("数据生命周期与前端工程化",
         "会话逻辑删除保留 90 天，@Scheduled(cron = \"0 0 3 * * *\") 每日凌晨物理清理并由外键 ON DELETE CASCADE "
         "级联删除消息；前端封装统一请求方法自动注入 Bearer 令牌，401 清理登录态并跳转，"
         "路由守卫同时校验登录态与 admin 角色。"),
    ]
    for label, text in bullets1:
        p = para(doc, before=2, line=LINE, left=18, first=-9)
        add_runs(p, [("• ", False), (label + "：", True), (text, False)])

    # ---------- 项目二 ----------
    p = para(doc, before=7.0, line=LINE)
    add_runs(p, [("RAG-Agent 智能客服问答系统  ", True), ("2026.4 - 2026.6", True)], NAME_SZ)

    p = para(doc, line=LINE_S)
    add_runs(p, [("github.com/JokerGHB/ai-customer-agent", False)], SMALL)

    p = para(doc, before=1, line=LINE_S)
    add_runs(p, [("技术栈：", True), (
        "FastAPI / LangChain / LangGraph / ChromaDB / BGE Reranker / 通义千问 / Redis / Docker / GitHub CI",
        False)], SMALL)

    p = para(doc, before=1.5, line=LINE)
    add_runs(p, [(
        "基于 RAG 知识库底座 + LangGraph 多工具 Agent 的智能问答 & 客服系统，覆盖文档预处理、检索增强、"
        "Agent 编排、安全防护、缓存限流与容器化部署全链路。", False)])

    bullets2 = [
        ("双阶段检索精排",
         "向量库粗召回 10 条，BM25 关键词检索补充实体召回，BGE Reranker 精排取 5 条，"
         "Hit Rate 提升 10%+，减少噪音片段进入 LLM 上下文。"),
        ("ReAct Agent 与安全防护",
         "编排 8 个业务工具，Checkpoint 持久化会话状态支持断点恢复；递归轮次上限、两级 Token 阈值、"
         "循环检测与三层校验中间件，工具调用准确率由 43% 提升至 91%。"),
        ("流式输出与会话记忆",
         "SSE 逐 token 推送，stream_id + offset 实现断线重连与断点续传；Checkpoint 与数据库分层记忆，"
         "支持跨会话复用用户关键信息。"),
        ("缓存限流与 CI 部署",
         "Redis 承载会话读缓存与路由缓存，IP 滑动窗口限流 30 次/分钟，TTL 随机抖动防雪崩；"
         "GitHub CI 完成静态检查、单元测试与 Docker 镜像构建。"),
    ]
    for label, text in bullets2:
        p = para(doc, before=2, line=LINE, left=18, first=-9)
        add_runs(p, [("• ", False), (label + "：", True), (text, False)])

    # ================= 技能清单 =================
    p = para(doc, before=7.5, after=7.0, line=LINE, size=HEAD)
    add_runs(p, [("技能清单", True)], HEAD)

    skills = [
        ("编程语言：", "Java / SQL / Python / JavaScript / TypeScript"),
        ("后端与框架：", "Spring Boot / Spring MVC / MyBatis / RESTful API / Maven / FastAPI"),
        ("数据与缓存：", "MySQL（索引设计、事务、预聚合）/ Redis / ChromaDB"),
        ("安全与工程：", "BCrypt / 令牌鉴权 / 统一异常处理 / 定时任务 / Docker / Git / GitHub CI"),
        ("前端能力：", "Vue 3 / Vite / ECharts / Axios / HTML + CSS"),
        ("AI/大模型：", "LangChain / LangGraph / RAG / Agent / MCP / Prompt Engineering"),
        ("开发工具：", "Cursor / Claude Code / Vibe Coding / AI 辅助编程"),
    ]
    for i, (label, value) in enumerate(skills):
        p = para(doc, before=2 if i else 0, line=LINE)
        add_runs(p, [(label, True), (value, False)])

    # ================= 荣誉证书 =================
    p = para(doc, before=7.5, after=7.0, line=LINE, size=HEAD)
    add_runs(p, [("荣誉证书", True)], HEAD)

    honors = [
        "蓝桥杯全国软件和信息技术专业人才大赛 — 省级二等奖",
        "校二等奖学金",
        "大学英语四级证书（CET-4）",
    ]
    for i, h in enumerate(honors):
        p = para(doc, before=2 if i else 0, line=LINE, left=18, first=-9)
        add_runs(p, [("• ", False), (h, False)])

    # ================= 分页控制 =================
    # 1) 所有段落禁止跨页拆分；2) 小节标题与项目抬头、链接、技术栈与其后内容同页
    heads = ("个人简介", "教育背景", "项目经历", "技能清单", "荣誉证书")
    for p in doc.paragraphs:
        txt = p.text.strip()
        p.paragraph_format.keep_together = True
        if not txt:
            continue
        if txt in heads:
            p.paragraph_format.keep_with_next = True
        elif txt.startswith("github.com") or txt.startswith("技术栈："):
            p.paragraph_format.keep_with_next = True
        elif (" 2026.10 - 至今" in txt) or (" 2026.4 - 2026.6" in txt):
            p.paragraph_format.keep_with_next = True

    doc.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    build()
